import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models import ConsultaCnpjStatus
from app.services.results import JobResults


EXCEL_MAX_CELL_LENGTH = 32_767
INVALID_EXCEL_CHARACTERS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
FORMULA_PREFIXES = ("=", "+", "-", "@")


def build_job_export(results: JobResults) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Resultados"
    worksheet.append(["Nome do CNPJ", "CNPJ", "Telefone", "E-mail"])

    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for item in results.items:
        if item.status == ConsultaCnpjStatus.SUCCESS.value:
            razao_social = _sanitize_text(item.razao_social)
            telefone = _sanitize_text(item.telefone)
            email = _sanitize_text(item.email)
        else:
            razao_social = None
            telefone = None
            email = None

        worksheet.append(
            [
                razao_social,
                _sanitize_text(item.cnpj),
                telefone,
                email,
            ]
        )
        row = worksheet.max_row
        worksheet.cell(row=row, column=2).number_format = "@"
        worksheet.cell(row=row, column=3).number_format = "@"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:D{worksheet.max_row}"
    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 16
    worksheet.column_dimensions["C"].width = 20
    worksheet.column_dimensions["D"].width = 40

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    buffer.seek(0)
    return buffer


def _sanitize_text(value: str | None) -> str | None:
    if value is None:
        return None

    sanitized = INVALID_EXCEL_CHARACTERS.sub("", value)[:EXCEL_MAX_CELL_LENGTH]
    if sanitized.startswith(FORMULA_PREFIXES):
        return f"'{sanitized}"
    return sanitized
