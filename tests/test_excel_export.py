from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

from app.services.excel_export import build_job_export
from app.services.results import JobResultItem, JobResults


def test_build_job_export_preserves_values_order_and_text_formatting() -> None:
    results = JobResults(
        id=uuid4(),
        status="COMPLETED",
        total=3,
        items=[
            JobResultItem("04252011000110", "SUCCESS", "Empresa Teste", "01123456789", None),
            JobResultItem("12345678000195", "NOT_FOUND", "Ignored", "Ignored", "Ignored"),
            JobResultItem("11444777000161", "FAILED", None, None, None),
        ],
    )

    workbook = load_workbook(BytesIO(build_job_export(results).getvalue()))
    worksheet = workbook.active

    assert [cell.value for cell in worksheet[1]] == ["Nome do CNPJ", "CNPJ", "Telefone", "E-mail"]
    assert list(worksheet.values)[1:] == [
        ("Empresa Teste", "04252011000110", "01123456789", None),
        (None, "12345678000195", None, None),
        (None, "11444777000161", None, None),
    ]
    assert worksheet["B2"].number_format == "@"
    assert worksheet["C2"].number_format == "@"
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref == "A1:D4"


def test_build_job_export_sanitizes_formula_prefixes_and_invalid_characters() -> None:
    results = JobResults(
        id=uuid4(),
        status="COMPLETED",
        total=1,
        items=[
            JobResultItem(
                "04252011000110",
                "SUCCESS",
                "=SUM(1,1)\x00",
                "+551123456789",
                "@example.com",
            )
        ],
    )

    workbook = load_workbook(BytesIO(build_job_export(results).getvalue()))
    worksheet = workbook.active

    assert worksheet["A2"].value == "'=SUM(1,1)"
    assert worksheet["C2"].value == "'+551123456789"
    assert worksheet["D2"].value == "'@example.com"
